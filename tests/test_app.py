"""Pruebas de la app completa con AppTest (modo local en una carpeta temporal)."""

from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

APP = str(Path(__file__).resolve().parent.parent / "app.py")


MASTER_PW = "master-pass-1"


@pytest.fixture
def at(tmp_path, monkeypatch):
    monkeypatch.setenv("HMCZ_LOCAL_DIR", str(tmp_path / "data"))
    import streamlit as st
    from core.auth import USERS_PATH, User, dump_users, hash_password
    from core.events import initialize
    from storage.local_store import LocalStore
    st.cache_resource.clear()
    st.cache_data.clear()
    store = LocalStore(tmp_path / "data")
    users = {"ana": User("ana", "Ana Admin", "admin", hash_password("admin-pass")),
             "sam": User("sam", "Sam Sup", "supervisor", hash_password("sup-pass1")),
             "old": User("old", "Old User", "supervisor", hash_password("old-pass1"), active=False)}
    store.write(USERS_PATH, dump_users(users), "test")
    initialize(store, "ana")
    app = AppTest.from_file(APP, default_timeout=60)
    app.secrets["master"] = {"username": "iibanez", "name": "Ivan", "password_hash": hash_password(MASTER_PW)}
    return app


def _login(at, username, password):
    at.run()
    at.text_input[0].input(username)
    at.text_input[1].input(password)
    at.button[0].click().run()
    return at


def _setup_admin(at):
    return _login(at, "ana", "admin-pass")


def test_login_always_shown_and_admin_view(at):
    at.run()
    assert "user" not in at.session_state
    assert len(at.text_input) == 2  # solo usuario y contraseña: no hay pantalla de alta
    _setup_admin(at)
    assert not at.exception
    assert len(at.tabs) == 10  # 7 pestañas + 3 de gestión de actividades (camión)
    assert any("Coordinate trunnion work" in m.value for m in at.markdown)
    assert not any(s.label == "View" for s in at.get("segmented_control"))  # el admin no ve la vista de supervisor


def test_public_view_shows_gantts_read_only(at, tmp_path):
    from core.events import write_event
    from storage.local_store import LocalStore
    write_event(LocalStore(tmp_path / "data"), "scope", {"op": "add_truck", "truck": "RD240", "acts": [
        {"act": "trailing_lower", "tasks": None}]}, "ana")
    at.run()
    assert not at.exception
    assert "user" not in at.session_state
    titles = [m.value for m in at.markdown]
    assert any("Gantt truck modification" in v for v in titles)
    assert any("Gantt operating machine" in v for v in titles)
    assert not at.tabs and not at.get("data_editor")  # nada editable sin iniciar sesión
    assert [b.label for b in at.button] == ["Sign in"]
    at.session_state["pub_mode"] = "activity"
    at.run()
    assert not at.exception


def test_login_fails_with_wrong_password(at):
    _login(at, "ana", "wrong-password")
    assert "user" not in at.session_state
    assert at.error


def test_inactive_user_cannot_login(at):
    _login(at, "old", "old-pass1")
    assert "user" not in at.session_state


def test_supervisor_sees_only_supervisor_view(at):
    _login(at, "sam", "sup-pass1")
    assert not at.exception
    assert len(at.tabs) == 0
    assert at.selectbox[0].options[0].startswith("RD")
    assert not any(s.label == "View" for s in at.get("segmented_control"))


def test_master_switches_views(at):
    _login(at, "iibanez", MASTER_PW)
    assert at.session_state["user"].is_master
    assert len(at.tabs) == 10  # 7 pestañas + 3 de gestión de actividades (camión)
    at.session_state["view_pick"] = "supervisor"
    at.run()
    assert not at.exception
    assert len(at.tabs) == 0


def test_master_in_users_file_is_ignored(at, tmp_path):
    from core.auth import USERS_PATH, User, dump_users, hash_password
    from storage.local_store import LocalStore
    LocalStore(tmp_path / "data").write(USERS_PATH, dump_users(
        {"iibanez": User("iibanez", "Fake", "admin", hash_password("fake-pass"))}), "test")
    _login(at, "iibanez", "fake-pass")
    assert "user" not in at.session_state


def test_japanese_ui(at):
    _setup_admin(at)
    at.session_state["lang_pick"] = "ja"
    at.run()
    assert not at.exception
    assert any("トラニオン" in m.value for m in at.markdown)


def test_admin_actions_shift_ack_settings(at, tmp_path):
    _setup_admin(at)
    truck = at.selectbox(key="plan_truck").value
    at.text_input(key=f"shift_reason_{truck}").input("Parts late")
    next(b for b in at.button if b.label == "Shift and save").click().run()
    assert not at.exception
    assert list((tmp_path / "data" / "log" / "plan").rglob("*.json"))
    at.text_input(key="ack_RD224").input("PO 123")
    at.button(key="ackb_RD224").click().run()
    assert list((tmp_path / "data" / "log" / "ack").rglob("*.json"))
    assert not any("RD224 · Coordinate" in m.value for m in at.markdown)


def test_supervisor_delay_note(at, tmp_path):
    _login(at, "sam", "sup-pass1")
    area = next(a for a in at.text_area if a.label == "Detailed description")
    area.input("Waiting for cover plates")
    next(b for b in at.button if b.label == "Submit delay note").click().run()
    assert not at.exception
    assert list((tmp_path / "data" / "log" / "delay").rglob("*.json"))


def test_excel_report(tmp_path):
    from datetime import date
    from openpyxl import load_workbook
    from core.events import initialize
    from core.project import build_project
    from core.reports import build_excel
    from storage.local_store import LocalStore

    s = LocalStore(tmp_path / "d")
    initialize(s, "x")
    data = build_excel(build_project(s.snapshot(), today=date(2026, 10, 25)))
    import io
    wb = load_workbook(io.BytesIO(data))
    assert {"Summary", "Trucks", "Activities", "SOP tasks", "Delay notes"} <= set(wb.sheetnames)
    assert wb["Trucks"].max_row == 13


def test_executive_pdf(tmp_path):
    from datetime import date
    from core.events import initialize
    from core.executive_report import build_executive_pdf
    from core.project import build_project
    from storage.local_store import LocalStore

    s = LocalStore(tmp_path / "d")
    initialize(s, "x")
    data = build_executive_pdf(build_project(s.snapshot(), today=date(2026, 10, 25)))
    assert data.startswith(b"%PDF") and len(data) > 10_000


def test_admin_truck_gantt_and_no_user_management(at):
    _setup_admin(at)
    assert at.session_state["g_mode"] == "truck"
    at.session_state["g_mode"] = "activity"
    at.run()
    assert not at.exception
    assert not any(b.label == "Add user" for b in at.button)
    # Los reportes solo se generan al presionar el botón
    assert not any(b.label == "Download executive report (PDF)" for b in at.get("download_button"))
    next(b for b in at.button if b.label == "Generate reports").click().run()
    assert not at.exception
    assert any(b.label == "Download executive report (PDF)" for b in at.get("download_button"))


def test_admin_edits_progress_and_scope(at, tmp_path):
    _setup_admin(at)
    truck = at.selectbox(key="det_truck").value
    # Avance manual de una actividad sin SOP desde la pestaña de camiones
    acts = at.selectbox(key=f"adm_act_{truck}").options
    at.selectbox(key=f"adm_act_{truck}").set_value(next(o for o in at.selectbox(key=f"adm_act_{truck}").options
                                                         if "Trailing arm lowering" in o)).run()
    assert not at.exception
    next(b for b in at.button if b.label == "Save progress").click().run()
    assert not at.exception
    assert list((tmp_path / "data" / "log" / "task").rglob("*.json"))
    assert acts
    # Nueva actividad general en todos los camiones
    at.text_input(key=f"new_name_{truck}_0").input("Final cleaning")
    at.toggle(key=f"new_all_{truck}_0").set_value(True)
    next(b for b in at.button if b.label == "Add activity").click().run()
    assert not at.exception
    assert list((tmp_path / "data" / "log" / "scope").rglob("*.json"))
    assert any("Final cleaning" in o for o in at.selectbox(key=f"adm_act_{truck}").options)
    # Camión nuevo
    at.text_input(key="nt_id_0").input("rd240")
    next(b for b in at.button if b.label == "Add truck").click().run()
    assert not at.exception
    assert any(o.startswith("RD240") for o in at.selectbox(key="det_truck").options)
    # Eliminar el camión recién agregado
    at.selectbox(key="det_truck").set_value("RD240").run()
    at.text_input(key="rt_id_RD240_0").input("RD240").run()
    next(b for b in at.button if b.label == "Remove truck").click().run()
    assert not at.exception
    assert not any(o.startswith("RD240") for o in at.selectbox(key="det_truck").options)
